import io

import pytest
from docx import Document
from pypdf import PdfWriter

from bot.services.extract import ExtractError, extract_text


def make_docx() -> bytes:
    d = Document()
    d.add_paragraph("Первый абзац.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "яч1"
    t.rows[0].cells[1].text = "яч2"
    d.add_paragraph("Последний абзац.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offs)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objs) + 1, xref)
    return out


def test_txt():
    assert extract_text("a.TXT", "Привет".encode("cp1251")) == "Привет"


def test_docx_paragraphs_and_tables_in_order():
    text = extract_text("a.docx", make_docx())
    assert text.index("Первый") < text.index("яч1 | яч2") < text.index("Последний")


def test_pdf_text():
    assert "Hello book" in extract_text("a.pdf", make_pdf("Hello book"))


def test_pdf_scan_without_text():
    w = PdfWriter()
    w.add_blank_page(200, 200)
    buf = io.BytesIO()
    w.write(buf)
    with pytest.raises(ExtractError, match="скан"):
        extract_text("a.pdf", buf.getvalue())


@pytest.mark.parametrize("name", ["a.docx", "a.pdf"])
def test_corrupt_files(name):
    with pytest.raises(ExtractError):
        extract_text(name, b"not a real file")


def test_unsupported():
    with pytest.raises(ExtractError):
        extract_text("a.doc", b"x")
